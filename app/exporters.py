"""Export a generated TrainMind package to DOCX, PDF, and PPTX."""

from __future__ import annotations

import io
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.shared import RGBColor as DocxRGBColor
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Flowable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

LEVEL_TITLES = {"beginner": "Beginner", "intermediate": "Intermediate", "advanced": "Advanced"}
ARTIFACT_TITLES = {"quiz": "Quiz", "exercise": "Exercise", "session_plan": "90-Minute Session Plan"}
NAVY = "#172B4D"
BLUE = "#3563E9"
TEAL = "#0F9D8A"
PAPER = "#F5F7FB"
GOLD = "#D99A32"
MUTED = "#66758F"
WHITE = "#FFFFFF"
EXPORT_VERSION = 3


def _all_artifacts(package: dict):
    for level in ("beginner", "intermediate", "advanced"):
        data = package["levels"][level]
        for kind in ("quiz", "exercise", "session_plan"):
            yield level, kind, data[kind]


def _plain_markdown(line: str) -> str:
    line = re.sub(r"^\s*[-*]\s+", "• ", line)
    line = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
    line = re.sub(r"`([^`]*)`", r"\1", line)
    return line.strip()


def _content_sections(content: str) -> list[tuple[str, list[str]]]:
    sections: list[tuple[str, list[str]]] = []
    heading = "Overview"
    lines: list[str] = []
    for line in content.splitlines():
        stripped = line.strip()
        match = re.match(r"^#{1,3}\s+(.+?)\s*$", stripped)
        if match:
            if lines or heading != "Overview":
                sections.append((heading, lines))
            heading = match.group(1).strip()
            lines = []
        elif stripped:
            lines.append(stripped)
    if lines or heading != "Overview":
        sections.append((heading, lines))
    return sections or [("Overview", ["No content generated."])]


def _rgb(hex_color: str) -> RGBColor:
    value = hex_color.lstrip("#")
    return RGBColor(*(int(value[index:index + 2], 16) for index in (0, 2, 4)))


def _add_textbox(slide, left, top, width, height, text, *, size, color, bold=False, font="Aptos"):
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = PptInches(0.02)
    frame.margin_right = PptInches(0.02)
    frame.margin_top = PptInches(0.02)
    frame.margin_bottom = PptInches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = font
    paragraph.font.size = PptPt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = _rgb(color)
    return shape


def _add_ppt_cover(prs: Presentation, topic: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = _rgb(NAVY)

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, PptInches(0), PptInches(0), PptInches(0.22), PptInches(7.5)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = _rgb(TEAL)
    accent.line.fill.background()

    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, PptInches(10.8), PptInches(-1.15), PptInches(4.2), PptInches(4.2)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = _rgb(BLUE)
    circle.fill.transparency = 15
    circle.line.fill.background()
    small_circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, PptInches(11.9), PptInches(4.9), PptInches(1.5), PptInches(1.5)
    )
    small_circle.fill.solid()
    small_circle.fill.fore_color.rgb = _rgb(TEAL)
    small_circle.fill.transparency = 10
    small_circle.line.fill.background()

    _add_textbox(slide, PptInches(0.9), PptInches(1.25), PptInches(8.5), PptInches(0.5),
                 "TRAINMIND  /  TRAINER COPILOT", size=15, color="#73D9CB", bold=True)
    _add_textbox(slide, PptInches(0.9), PptInches(2.15), PptInches(10.8), PptInches(1.35),
                 topic, size=36, color=WHITE, bold=True)
    _add_textbox(slide, PptInches(0.92), PptInches(3.7), PptInches(9.8), PptInches(0.65),
                 subtitle, size=20, color="#D9E3F3")
    _add_textbox(slide, PptInches(0.92), PptInches(5.55), PptInches(8.8), PptInches(0.75),
                 "Curriculum-grounded learning materials\nQuiz  ·  Exercise  ·  Session Plan",
                 size=15, color="#C5D1E4")


def _add_ppt_content_slide(
    prs: Presentation,
    topic: str,
    level: str,
    kind: str,
    section_title: str,
    lines: list[str],
    section_number: int,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(PAPER)

    top_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, PptInches(0), PptInches(0), PptInches(13.333), PptInches(0.16)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = _rgb(TEAL if kind == "session_plan" else BLUE)
    top_bar.line.fill.background()

    _add_textbox(slide, PptInches(0.68), PptInches(0.42), PptInches(9.4), PptInches(0.28),
                 f"{LEVEL_TITLES[level].upper()}  /  {ARTIFACT_TITLES[kind].upper()}",
                 size=10, color=TEAL if kind == "session_plan" else BLUE, bold=True)
    _add_textbox(slide, PptInches(0.68), PptInches(0.82), PptInches(11.9), PptInches(0.62),
                 section_title, size=25, color=NAVY, bold=True)

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        PptInches(0.68), PptInches(1.72), PptInches(11.95), PptInches(4.95),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = _rgb(WHITE)
    card.line.color.rgb = _rgb("#E2E8F2")
    card.line.width = PptPt(1)

    content = slide.shapes.add_textbox(
        PptInches(1.02), PptInches(2.02), PptInches(11.25), PptInches(4.35)
    )
    frame = content.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = PptInches(0.04)
    frame.margin_right = PptInches(0.04)
    frame.margin_top = PptInches(0.04)
    frame.margin_bottom = PptInches(0.04)

    for index, line in enumerate(lines or ["No content generated."]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        clean = _plain_markdown(line)
        is_bullet = line.startswith(("-", "*"))
        paragraph.text = f"•  {clean}" if is_bullet else clean
        paragraph.font.name = "Aptos"
        paragraph.font.size = PptPt(15 if len(clean) < 150 else 12)
        paragraph.font.color.rgb = _rgb(MUTED if clean.lower().startswith(("**curriculum evidence", "curriculum evidence")) else NAVY)
        paragraph.space_after = PptPt(11 if is_bullet else 8)
        if re.match(r"^\*\*[^*]+:\*\*", line):
            paragraph.font.bold = True
            paragraph.font.color.rgb = _rgb(TEAL)

    _add_textbox(slide, PptInches(0.7), PptInches(6.95), PptInches(10.8), PptInches(0.25),
                 topic, size=9, color=MUTED)
    _add_textbox(slide, PptInches(11.25), PptInches(6.95), PptInches(1.25), PptInches(0.25),
                 f"{section_number:02d}", size=9, color=MUTED, bold=True)


def _export_styled_pptx(topic: str, artifacts: list[tuple[str, str, str]]) -> bytes:
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    subtitle = "Complete training package" if len(artifacts) > 1 else (
        f"{LEVEL_TITLES[artifacts[0][0]]}  /  {ARTIFACT_TITLES[artifacts[0][1]]}"
    )
    _add_ppt_cover(prs, topic, subtitle)

    for level, kind, content in artifacts:
        sections = _content_sections(content)
        if kind == "quiz":
            question_sections = [section for section in sections if section[0].lower().startswith("question")]
            display_sections = question_sections or sections
        elif kind == "exercise":
            exercise_sections = [
                section for section in sections
                if section[0].lower() not in {"exercise", "retrieved curriculum sources"}
            ]
            display_sections = []
            for start in range(0, len(exercise_sections), 3):
                group = exercise_sections[start:start + 3]
                group_lines = []
                for title, section_lines in group:
                    group_lines.append(f"**{title}:**")
                    group_lines.extend(section_lines)
                display_sections.append((
                    "Exercise Details" if start == 0 else "Exercise Details  /  continued",
                    group_lines,
                ))
        else:
            display_sections = [section for section in sections if section[0].lower() not in {
                "quiz", "exercise", "90-minute session schedule", "retrieved curriculum sources",
            }]
            if kind == "session_plan":
                schedule = [section for section in sections if re.match(r"\d+\s*[-–]\s*\d+\s*min", section[0], re.I)]
                intro = [section for section in sections if section[0].lower() in {"learning objectives", "expected learning outcome", "curriculum support note"}]
                display_sections = intro + schedule
        if not display_sections:
            display_sections = sections

        for section_title, section_lines in display_sections:
            chunks = [section_lines[index:index + 9] for index in range(0, len(section_lines), 9)] or [[]]
            for chunk_index, chunk in enumerate(chunks, start=1):
                title = section_title
                if len(chunks) > 1:
                    title += f"  /  continued {chunk_index}"
                _add_ppt_content_slide(
                    prs, topic, level, kind, title, chunk, len(prs.slides) + 1
                )

    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def _export_docx(topic: str, artifacts: list[tuple[str, str, str]]) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    for style_name, color, size in (
        ("Title", NAVY, 30),
        ("Heading 1", NAVY, 21),
        ("Heading 2", BLUE, 14),
        ("Heading 3", TEAL, 11),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name in {"Title", "Heading 1"} else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = DocxRGBColor.from_string(color.lstrip("#"))
        style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    body_style = styles["Normal"]
    body_style.font.name = "Aptos"
    body_style.font.size = Pt(10)
    body_style.font.color.rgb = DocxRGBColor.from_string(NAVY.lstrip("#"))
    body_style.paragraph_format.space_after = Pt(5)
    body_style.paragraph_format.line_spacing = 1.12

    header = section.header.paragraphs[0]
    header.text = "TRAINMIND  /  TRAINER COPILOT"
    header.style = "Caption"
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = DocxRGBColor.from_string(TEAL.lstrip("#"))
    header.paragraph_format.space_after = Pt(5)
    _set_paragraph_bottom_border(header, "DCE4F0")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("TRAINMIND  ·  ")
    _add_docx_page_number(footer)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = DocxRGBColor.from_string(MUTED.lstrip("#"))

    cover_table = doc.add_table(rows=1, cols=1)
    cover_table.autofit = False
    cover_cell = cover_table.cell(0, 0)
    cover_cell.width = Inches(6.85)
    cover_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_shading(cover_cell, NAVY)
    _set_cell_margins(cover_cell, top=950, start=650, bottom=950, end=650)
    brand = cover_cell.paragraphs[0]
    brand.paragraph_format.space_after = Pt(20)
    brand_run = brand.add_run("TRAINMIND  /  TRAINER COPILOT")
    brand_run.bold = True
    brand_run.font.size = Pt(11)
    brand_run.font.color.rgb = DocxRGBColor.from_string("73D9CB")
    cover_title = cover_cell.add_paragraph()
    cover_title.paragraph_format.space_after = Pt(12)
    title_run = cover_title.add_run(topic)
    title_run.bold = True
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(30)
    title_run.font.color.rgb = DocxRGBColor.from_string(WHITE.lstrip("#"))
    cover_subtitle = cover_cell.add_paragraph("Curriculum-grounded learning materials")
    cover_subtitle.paragraph_format.space_after = Pt(6)
    for run in cover_subtitle.runs:
        run.font.size = Pt(13)
        run.font.color.rgb = DocxRGBColor.from_string("D9E3F3")
    deliverables = cover_cell.add_paragraph("QUIZ   ·   EXERCISE   ·   SESSION PLAN")
    for run in deliverables.runs:
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = DocxRGBColor.from_string("73D9CB")
    doc.add_page_break()

    for index, (level, kind, content) in enumerate(artifacts):
        artifact_heading = doc.add_heading(level=1)
        artifact_run = artifact_heading.add_run(
            f"{LEVEL_TITLES[level]}   /   {ARTIFACT_TITLES[kind]}"
        )
        artifact_run.font.color.rgb = DocxRGBColor.from_string(
            (TEAL if kind == "session_plan" else BLUE).lstrip("#")
        )
        _set_paragraph_bottom_border(artifact_heading, TEAL if kind == "session_plan" else BLUE)

        for line in content.splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("### "):
                doc.add_heading(_plain_markdown(stripped[4:]), level=3)
            elif stripped.startswith("## "):
                doc.add_heading(_plain_markdown(stripped[3:]), level=2)
            elif stripped.startswith("# "):
                doc.add_heading(_plain_markdown(stripped[2:]), level=1)
            elif stripped.startswith("- ") or stripped.startswith("* "):
                paragraph = doc.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.left_indent = Inches(0.22)
                paragraph.paragraph_format.first_line_indent = Inches(-0.14)
                run = paragraph.add_run(_plain_markdown(stripped))
                run.font.color.rgb = DocxRGBColor.from_string(NAVY.lstrip("#"))
            else:
                label_match = re.match(
                    r"^\*\*(Correct answer|Explanation|Curriculum evidence|Trainer actions|Learner actions|Materials from curriculum):\*\*\s*(.*)$",
                    stripped,
                    re.IGNORECASE,
                )
                if label_match:
                    paragraph = doc.add_paragraph()
                    paragraph.paragraph_format.left_indent = Inches(0.12)
                    paragraph.paragraph_format.right_indent = Inches(0.08)
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.space_after = Pt(7)
                    _set_paragraph_shading(paragraph, "#EAF5F4")
                    label = paragraph.add_run(f"{label_match.group(1)}: ")
                    label.bold = True
                    label.font.color.rgb = DocxRGBColor.from_string(TEAL.lstrip("#"))
                    paragraph.add_run(_plain_markdown(label_match.group(2)))
                else:
                    doc.add_paragraph(_plain_markdown(stripped))
        if index < len(artifacts) - 1:
            doc.add_page_break()

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_margins(cell, *, top: int, start: int, bottom: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def _set_paragraph_shading(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill.lstrip("#"))
    properties.append(shading)


def _set_paragraph_bottom_border(paragraph, color: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color.lstrip("#"))
    borders.append(bottom)
    properties.append(borders)


def _add_docx_page_number(paragraph) -> None:
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def export_docx(package: dict) -> bytes:
    return _export_docx(package["topic"], list(_all_artifacts(package)))


class _CoverPage(Flowable):
    def __init__(self, topic: str, subtitle: str):
        super().__init__()
        self.topic = topic
        self.subtitle = subtitle
        self.width = A4[0]
        self.height = A4[1] - 40 * mm

    def draw(self):
        canvas = self.canv
        canvas.saveState()
        canvas.setFillColor(colors.HexColor(NAVY))
        canvas.roundRect(0, 0, self.width, self.height, 6 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.HexColor(TEAL))
        canvas.roundRect(0, 0, 5 * mm, self.height, 2 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.HexColor("#73D9CB"))
        canvas.setFont("Helvetica-Bold", 10)
        canvas.drawString(22 * mm, 232 * mm, "TRAINMIND  /  TRAINER COPILOT")
        canvas.setFillColor(colors.white)
        topic_lines = _wrap_text(self.topic, 38)
        canvas.setFont("Helvetica-Bold", 27)
        y = 195 * mm
        for line in topic_lines[:4]:
            canvas.drawString(22 * mm, y, line)
            y -= 13 * mm
        canvas.setFillColor(colors.HexColor("#D9E3F3"))
        canvas.setFont("Helvetica", 14)
        canvas.drawString(22 * mm, y - 4 * mm, self.subtitle)
        canvas.setFillColor(colors.HexColor("#C5D1E4"))
        canvas.setFont("Helvetica", 10)
        canvas.drawString(22 * mm, 66 * mm, "Curriculum-grounded learning materials")
        canvas.restoreState()


def _wrap_text(text: str, width: int) -> list[str]:
    words = text.split()
    lines = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if len(candidate) > width and current:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def _pdf_header_footer(canvas, _doc):
    canvas.saveState()
    page = canvas.getPageNumber()
    if page > 1:
        canvas.setStrokeColor(colors.HexColor("#DCE4F0"))
        canvas.setLineWidth(0.7)
        canvas.line(18 * mm, A4[1] - 13 * mm, A4[0] - 18 * mm, A4[1] - 13 * mm)
        canvas.setFont("Helvetica-Bold", 8)
        canvas.setFillColor(colors.HexColor(NAVY))
        canvas.drawString(18 * mm, A4[1] - 10 * mm, "TRAINMIND")
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"{page - 1}")
    canvas.restoreState()


def _pdf_markup(text: str) -> str:
    escaped = escape(_plain_markdown(text))
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)


def _pdf_section_band(title: str, kind: str):
    background = colors.HexColor(TEAL if kind == "session_plan" else BLUE)
    heading = Paragraph(
        f'<font color="#FFFFFF"><b>{escape(title)}</b></font>',
        ParagraphStyle(
            f"Band-{kind}",
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=15,
            textColor=colors.white,
        ),
    )
    band = Table([[heading]], colWidths=[174 * mm])
    band.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), background),
        ("BOX", (0, 0), (-1, -1), 0, background),
        ("LEFTPADDING", (0, 0), (-1, -1), 9),
        ("RIGHTPADDING", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    return band


def _export_pdf(topic: str, artifacts: list[tuple[str, str, str]]) -> bytes:
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm,
        topMargin=17 * mm, bottomMargin=17 * mm,
        title=f"TrainMind Training Package - {topic}",
    )
    base = getSampleStyleSheet()
    h2 = ParagraphStyle("PackageH2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=colors.HexColor(BLUE), spaceBefore=9, spaceAfter=5)
    body = ParagraphStyle("PackageBody", parent=base["BodyText"], fontName="Helvetica", fontSize=9.5, leading=14, textColor=colors.HexColor(NAVY), spaceAfter=5, wordWrap="CJK")
    bullet = ParagraphStyle("PackageBullet", parent=body, leftIndent=11, firstLineIndent=-8, bulletIndent=0, spaceAfter=4)
    callout = ParagraphStyle("PackageCallout", parent=body, textColor=colors.HexColor(NAVY), backColor=colors.HexColor("#EAF5F4"), borderColor=colors.HexColor("#D0EBE7"), borderWidth=0.5, borderPadding=7, leftIndent=6, rightIndent=6, spaceBefore=3, spaceAfter=7)
    subtitle = (
        "Complete curriculum-grounded training package"
        if len(artifacts) > 1
        else f"{LEVEL_TITLES[artifacts[0][0]]}  /  {ARTIFACT_TITLES[artifacts[0][1]]}"
    )
    story = [_CoverPage(topic, subtitle), PageBreak()]

    for index, (level, kind, content) in enumerate(artifacts):
        story.append(_pdf_section_band(
            f"{LEVEL_TITLES[level]}   /   {ARTIFACT_TITLES[kind]}", kind
        ))
        story.append(Spacer(1, 7 * mm))
        sections = _content_sections(content)
        for section_title, lines in sections:
            if section_title.lower() in {"quiz", "exercise", "90-minute session schedule"}:
                continue
            story.append(Paragraph(escape(section_title), h2))
            for line in lines:
                if line.startswith(("-", "*")):
                    story.append(Paragraph(_pdf_markup(line), bullet, bulletText="•"))
                elif line.lower().startswith(("**correct answer", "**explanation", "**curriculum evidence", "**trainer actions", "**learner actions", "**materials from curriculum")):
                    story.append(Paragraph(_pdf_markup(line), callout))
                else:
                    story.append(Paragraph(_pdf_markup(line), body))
        if index < len(artifacts) - 1:
            story.append(PageBreak())

    doc.build(story, onFirstPage=_pdf_header_footer, onLaterPages=_pdf_header_footer)
    return output.getvalue()


def export_pdf(package: dict) -> bytes:
    return _export_pdf(package["topic"], list(_all_artifacts(package)))


def _export_pptx(topic: str, artifacts: list[tuple[str, str, str]]) -> bytes:
    return _export_styled_pptx(topic, artifacts)


def export_pptx(package: dict) -> bytes:
    return _export_pptx(package["topic"], list(_all_artifacts(package)))


def build_artifact_exports(topic: str, level: str, kind: str, content: str) -> dict[str, bytes]:
    """Build PDF, PowerPoint, and Word files for one individual artifact."""
    artifact = [(level, kind, content)]
    return {
        "pdf": _export_pdf(topic, artifact),
        "pptx": _export_pptx(topic, artifact),
        "docx": _export_docx(topic, artifact),
    }


def build_exports(package: dict) -> dict:
    """Build full-package exports and individual exports for every artifact."""
    return {
        "version": EXPORT_VERSION,
        "pdf": export_pdf(package),
        "pptx": export_pptx(package),
        "docx": export_docx(package),
        "artifacts": {
            level: {
                kind: build_artifact_exports(package["topic"], level, kind, content)
                for kind, content in package["levels"][level].items()
                if kind in ARTIFACT_TITLES
            }
            for level in ("beginner", "intermediate", "advanced")
        },
    }
