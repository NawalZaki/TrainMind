from io import BytesIO

import pymupdf
from docx import Document
from pptx import Presentation

from app.exporters import build_exports
from app.exporters import EXPORT_VERSION


def test_exporters_create_pdf_powerpoint_and_word():
    package = {"topic": "RAG", "levels": {}}
    for level in ("beginner", "intermediate", "advanced"):
        package["levels"][level] = {
            "quiz": (
                "## Quiz\n### Question 1\nWhich workflow is described?\n"
                "A) Retrieval\nB) Storage\nC) Training\nD) None\n"
                "**Correct answer:** A\n"
                "**Curriculum evidence:** sample.pdf, Page 1"
            ),
            "exercise": (
                "## Exercise\n## Objective\nExplain the workflow.\n"
                "## Task\nUse the provided curriculum.\n"
                "## Steps\n- Read the source\n- Describe the process\n"
                "## Expected Outcome\nA source-grounded explanation.\n"
                "## Retrieved Curriculum Sources\n- sample.pdf, page 1"
            ),
            "session_plan": (
                "## Learning Objectives\n- Explain the workflow\n- Identify its steps\n"
                "## 90-Minute Session Schedule\n"
                "### 0-10 min: Introduction\n**Trainer actions:** Introduce topic.\n"
                "**Learner actions:** Share prior knowledge.\n"
                "### 10-25 min: Concept Explanation\n"
                "**Trainer actions:** Explain the workflow.\n"
                "**Learner actions:** Identify key steps.\n"
                "## Retrieved Curriculum Sources\n- sample.pdf, page 1"
            ),
            "timing_validation": {"valid": True, "total_minutes": 90},
        }

    files = build_exports(package)

    assert files["version"] == EXPORT_VERSION
    assert files["pdf"].startswith(b"%PDF")
    assert files["pptx"].startswith(b"PK")
    assert files["docx"].startswith(b"PK")
    assert all(len(files[extension]) > 1000 for extension in ("pdf", "pptx", "docx"))

    slides = Presentation(BytesIO(files["pptx"])).slides
    assert len(slides) >= 7
    cover_text = " ".join(shape.text for shape in slides[0].shapes if shape.has_text_frame)
    assert "RAG" in cover_text
    assert "TRAINMIND" in cover_text
    assert slides[0].background.fill.fore_color.rgb == (23, 43, 77)
    all_slide_text = " ".join(
        shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "Question 1" in all_slide_text
    assert "Exercise Details" in all_slide_text
    assert "Introduction" in all_slide_text

    pdf = pymupdf.open(stream=files["pdf"], filetype="pdf")
    assert len(pdf) > 2
    assert "RAG" in pdf[0].get_text()
    assert "TRAINMIND" in pdf[0].get_text()
    all_pdf_text = " ".join(page.get_text() for page in pdf)
    assert "Correct answer" in all_pdf_text
    assert "sample.pdf" in all_pdf_text
    pdf.close()

    word = Document(BytesIO(files["docx"]))
    assert len(word.tables) == 1
    cover_text = " ".join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
    assert "RAG" in cover_text
    assert "SESSION PLAN" in cover_text
    assert word.sections[0].header.paragraphs[0].text.startswith("TRAINMIND")
    all_word_text = " ".join(paragraph.text for paragraph in word.paragraphs)
    assert "Question 1" in all_word_text
    assert "sample.pdf" in all_word_text

    for level in ("beginner", "intermediate", "advanced"):
        for artifact in ("quiz", "exercise", "session_plan"):
            individual_exports = files["artifacts"][level][artifact]
            assert individual_exports["pdf"].startswith(b"%PDF")
            assert individual_exports["pptx"].startswith(b"PK")
            assert individual_exports["docx"].startswith(b"PK")
            artifact_slides = Presentation(BytesIO(individual_exports["pptx"])).slides
            assert len(artifact_slides) > 1
            artifact_pdf = pymupdf.open(stream=individual_exports["pdf"], filetype="pdf")
            assert len(artifact_pdf) == 2
            artifact_pdf.close()
